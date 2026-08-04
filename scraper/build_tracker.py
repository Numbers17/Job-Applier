"""Generates an application-tracking spreadsheet
(output/application_tracker.xlsx) from output/candidates.csv -- a
persistent, human-friendly record of what you've applied to, skipped, or
still need to review. Re-run any time; it always reflects the current
state of candidates.csv, which autofill.py updates as you mark rows.
"""
import argparse
import csv
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CANDIDATES_PATH = ROOT / "output" / "candidates.csv"
TRACKER_PATH = ROOT / "output" / "application_tracker.xlsx"
PROFILE_PATH = ROOT / "profile.json"
TRACKER_FILENAME = "application_tracker.xlsx"


def drive_sync_path() -> Path | None:
    """Optional second save location -- a Google Drive Desktop (Mirror mode)
    folder, set via profile.json's "drive_sync_path". Writing here directly
    lets Drive's own sync client pick up the change; no API/OAuth needed.
    """
    if not PROFILE_PATH.exists():
        return None
    profile = json.loads(PROFILE_PATH.read_text(encoding="utf-8"))
    raw_path = profile.get("drive_sync_path")
    return Path(raw_path) if raw_path else None

COLUMNS = [
    "Company",
    "Role",
    "Type",
    "Location",
    "Status",
    "Date Posted",
    "Status Updated",
    "ATS",
    "Job URL",
    "Source",
    "Notes",
]
WIDE_COLUMNS = {"Role": 40, "Job URL": 50, "Notes": 40, "Location": 24, "Source": 34}
TYPE_LABELS = {"internship": "Internship", "new_grad": "New Grad"}


def load_rows() -> list[dict]:
    if not CANDIDATES_PATH.exists():
        return []
    with CANDIDATES_PATH.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def display_status(row: dict) -> str:
    return row.get("status", "").strip() or "Not Started"


def build(only_reviewed: bool) -> None:
    rows = load_rows()
    if not rows:
        print(f"No data in {CANDIDATES_PATH} -- run run_pipeline.py first.")
        return
    if only_reviewed:
        rows = [r for r in rows if r.get("status", "").strip()]

    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"

    counts = Counter(display_status(r) for r in rows)
    ws.append(["Status", "Count"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for status, count in sorted(counts.items()):
        ws.append([status, count])
    ws.append(["Total", len(rows)])
    ws.append([])

    type_counts = Counter(TYPE_LABELS.get(r.get("listing_type", ""), r.get("listing_type", "") or "Unknown") for r in rows)
    ws.append(["Type", "Count"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for listing_type, count in sorted(type_counts.items()):
        ws.append([listing_type, count])
    ws.append([])

    header_row_idx = ws.max_row + 1
    ws.append(COLUMNS)
    for cell in ws[header_row_idx]:
        cell.font = Font(bold=True)
    ws.freeze_panes = f"A{header_row_idx + 1}"

    rows_sorted = sorted(rows, key=lambda r: r.get("status_updated", ""), reverse=True)
    url_col = COLUMNS.index("Job URL") + 1
    for r in rows_sorted:
        ws.append(
            [
                r.get("company", ""),
                r.get("role", ""),
                TYPE_LABELS.get(r.get("listing_type", ""), r.get("listing_type", "")),
                r.get("location", ""),
                display_status(r),
                r.get("date_posted", ""),
                r.get("status_updated", ""),
                r.get("ats", ""),
                r.get("url", ""),
                r.get("source", ""),
                "",  # Notes -- yours to fill in
            ]
        )
        url_cell = ws.cell(row=ws.max_row, column=url_col)
        if r.get("url"):
            url_cell.hyperlink = r["url"]
            url_cell.font = Font(color="0563C1", underline="single")

    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDE_COLUMNS.get(col, max(12, len(col) + 4))

    TRACKER_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(TRACKER_PATH)
    print(f"Wrote {len(rows)} rows ({dict(counts)}) to {TRACKER_PATH.relative_to(ROOT)}")

    sync_dir = drive_sync_path()
    if sync_dir:
        try:
            sync_dir.mkdir(parents=True, exist_ok=True)
            wb.save(sync_dir / TRACKER_FILENAME)
            print(f"Synced to {sync_dir / TRACKER_FILENAME}")
        except OSError as exc:
            print(f"  [warn] couldn't write to Drive sync folder ({sync_dir}): {exc}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--only-reviewed",
        action="store_true",
        help="Only include rows you've actually marked a status for (skip 'Not Started' rows)",
    )
    args = parser.parse_args()
    build(args.only_reviewed)
